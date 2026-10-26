#!/usr/bin/env python3
"""Build docs/Technical_Report.docx (and .pdf) from docs/report/Technical_Report.md.

Run from the project folder:   python3 handoff/build_report.py

Steps: pandoc (markdown -> docx), style pass (A4, fonts, tables, footer page numbers),
render a PDF with LibreOffice, read the page of every heading from the PDF, then a
second style pass that writes the table of contents with those page numbers.
Needs: pandoc, python-docx, LibreOffice, pdftotext (poppler).
Edit the markdown, never the docx: the docx is rebuilt every time.
"""
import json, os, re, shutil, subprocess, sys, tempfile

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, 'docs', 'report', 'Technical_Report.md')
OUT_DOCX = os.path.join(ROOT, 'docs', 'Technical_Report.docx')
OUT_PDF = os.path.join(ROOT, 'docs', 'Technical_Report.pdf')
FONT = 'Times New Roman'
MONO = 'Courier New'
TEXT_W_CM = 21.0 - 2 * 2.54  # A4 width minus margins


# ---------------------------------------------------------------- helpers
def run(cmd, **kw):
    r = subprocess.run(cmd, capture_output=True, text=True, **kw)
    if r.returncode != 0:
        sys.exit('FAILED: %s\n%s\n%s' % (' '.join(cmd), r.stdout, r.stderr))
    return r


def set_font(style, name=FONT, size=None, bold=None, italic=None, color=None):
    f = style.font
    f.name = name
    if size is not None:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if color is not None:
        f.color.rgb = RGBColor.from_string(color)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    for a in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
        if rfonts.get(qn(a)) is not None:
            del rfonts.attrib[qn(a)]
    for a in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        rfonts.set(qn(a), name)
    col = rpr.find(qn('w:color'))
    if col is not None and color is not None:
        for a in ('w:themeColor', 'w:themeShade', 'w:themeTint'):
            if col.get(qn(a)) is not None:
                del col.attrib[qn(a)]


def get_style(doc, name):
    # pandoc names its styles 'Heading 1' (capital H); python-docx looks up 'heading 1',
    # so search by name without regard to case.
    for s in doc.styles:
        if s.name is not None and s.name.lower() == name.lower():
            return s
    return None


def add_field(paragraph, instr):
    def el(tag, **attrs):
        e = OxmlElement(tag)
        for k, v in attrs.items():
            e.set(qn(k.replace('_', ':')), v)
        return e
    r1 = paragraph.add_run(); r1._r.append(el('w:fldChar', w_fldCharType='begin'))
    r2 = paragraph.add_run(); t = el('w:instrText'); t.set(qn('xml:space'), 'preserve'); t.text = instr; r2._r.append(t)
    r3 = paragraph.add_run(); r3._r.append(el('w:fldChar', w_fldCharType='separate'))
    r4 = paragraph.add_run('1')
    r5 = paragraph.add_run(); r5._r.append(el('w:fldChar', w_fldCharType='end'))
    return [r1, r2, r3, r4, r5]


def set_cell_shading(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    for old in tcpr.findall(qn('w:shd')):
        tcpr.remove(old)
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), fill)
    tcpr.append(shd)


def set_table_borders(table):
    tblpr = table._tbl.tblPr
    for old in tblpr.findall(qn('w:tblBorders')):
        tblpr.remove(old)
    b = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement('w:' + side)
        e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:space'), '0'); e.set(qn('w:color'), '808080')
        b.append(e)
    tblpr.append(b)


def set_cell_margins(table, top=40, bottom=40, left=80, right=80):
    tblpr = table._tbl.tblPr
    for old in tblpr.findall(qn('w:tblCellMar')):
        tblpr.remove(old)
    m = OxmlElement('w:tblCellMar')
    for side, v in (('top', top), ('left', left), ('bottom', bottom), ('right', right)):
        e = OxmlElement('w:' + side); e.set(qn('w:w'), str(v)); e.set(qn('w:type'), 'dxa'); m.append(e)
    tblpr.append(m)


def column_weights(table):
    ncols = len(table.columns)
    w = [1.0] * ncols
    for row in table.rows:
        for i, cell in enumerate(row.cells[:ncols]):
            txt = cell.text.strip()
            longest_word = max([len(x) for x in txt.split()] or [1])
            w[i] = max(w[i], min(len(txt), 70) * 0.55 + longest_word * 0.45)
    return [max(x, 11.0) for x in w]


def layout_table(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tblpr = table._tbl.tblPr
    for old in tblpr.findall(qn('w:tblLayout')):
        tblpr.remove(old)
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblpr.append(lay)
    for old in tblpr.findall(qn('w:tblW')):
        tblpr.remove(old)
    tw = OxmlElement('w:tblW'); tw.set(qn('w:w'), str(int(TEXT_W_CM / 2.54 * 1440))); tw.set(qn('w:type'), 'dxa'); tblpr.append(tw)
    weights = column_weights(table)
    total = sum(weights)
    widths = [Cm(TEXT_W_CM * x / total) for x in weights]
    grid = table._tbl.tblGrid
    for i, gc in enumerate(grid.findall(qn('w:gridCol'))):
        gc.set(qn('w:w'), str(int(widths[i].twips)))
    for row in table.rows:
        for i, cell in enumerate(row.cells[:len(widths)]):
            cell.width = widths[i]


# ---------------------------------------------------------------- styling
def style_document(path_in, path_out, pages=None):
    doc = Document(path_in)
    st = doc.styles

    # page: A4, 2.54 cm margins, no page number on the cover
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Cm(2.54)
    sec.header_distance = sec.footer_distance = Cm(1.25)
    sec.gutter = Cm(0)
    sec.different_first_page_header_footer = True
    fp = sec.footer.paragraphs[0] if sec.footer.paragraphs else sec.footer.add_paragraph()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in add_field(fp, ' PAGE '):
        r.font.name = FONT; r.font.size = Pt(10)

    # base text
    for name in ('Normal', 'Body Text', 'First Paragraph', 'Compact'):
        s = get_style(doc, name)
        if s is not None:
            set_font(s, size=12 if name != 'Compact' else 11)
            pf = s.paragraph_format
            pf.line_spacing = 1.15
            pf.space_after = Pt(6) if name != 'Compact' else Pt(2)
            pf.space_before = Pt(0)
    for name, size, before, after in (('Heading 1', 18, 0, 12), ('Heading 2', 14, 14, 6), ('Heading 3', 12, 10, 4)):
        s = get_style(doc, name)
        if s is not None:
            set_font(s, size=size, bold=True, italic=False, color='000000')
            pf = s.paragraph_format
            pf.space_before = Pt(before); pf.space_after = Pt(after)
            pf.keep_with_next = True
            pf.page_break_before = (name == 'Heading 1')
    s = get_style(doc, 'Verbatim Char')
    if s is not None:
        set_font(s, name=MONO, size=10)
    for name in ('Image Caption', 'Table Caption'):
        s = get_style(doc, name)
        if s is not None:
            set_font(s, size=10, italic=False, bold=(name == 'Table Caption'))
            s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            s.paragraph_format.space_after = Pt(10)
            s.paragraph_format.space_before = Pt(4)
    tc = get_style(doc, 'Table Caption')
    if tc is not None:
        tc.paragraph_format.keep_with_next = True
    cf = get_style(doc, 'Captioned Figure')
    if cf is not None:
        cf.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cf.paragraph_format.keep_with_next = True
        cf.paragraph_format.space_before = Pt(8)

    # cover and contents styles (made by pandoc from custom-style)
    cover = {'CoverSmall': (14, False, 14), 'CoverTitle': (24, True, 36), 'CoverNames': (14, False, 12)}
    for name, (size, bold, after) in cover.items():
        s = get_style(doc, name)
        if s is not None:
            set_font(s, size=size, bold=bold)
            s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            s.paragraph_format.space_after = Pt(after)
            s.paragraph_format.space_before = Pt(40 if name == 'CoverTitle' else 0)
    s = get_style(doc, 'TocTitle')
    if s is not None:
        set_font(s, size=18, bold=True)
        s.paragraph_format.page_break_before = True
        s.paragraph_format.space_after = Pt(14)

    # tables
    for table in doc.tables:
        set_table_borders(table)
        set_cell_margins(table)
        layout_table(table)
        for ri, row in enumerate(table.rows):
            trpr = row._tr.get_or_add_trPr()
            cant = OxmlElement('w:cantSplit'); trpr.append(cant)
            if ri == 0:
                hdr = OxmlElement('w:tblHeader'); trpr.append(hdr)
            for cell in row.cells:
                if ri == 0:
                    set_cell_shading(cell, 'E7E6E6')
                for para in cell.paragraphs:
                    para.paragraph_format.space_after = Pt(1)
                    para.paragraph_format.line_spacing = 1.0
                    for r in para.runs:
                        if r.font.name != MONO and 'Verbatim' not in (r.style.name if r.style is not None else ''):
                            r.font.size = Pt(10)
                        else:
                            r.font.size = Pt(9)
                        if ri == 0:
                            r.font.bold = True

    # put the children of every tblPr into the order the schema requires
    order = ['tblStyle', 'tblpPr', 'tblOverlap', 'bidiVisual', 'tblStyleRowBandSize', 'tblStyleColBandSize',
             'tblW', 'jc', 'tblCellSpacing', 'tblInd', 'tblBorders', 'shd', 'tblLayout', 'tblCellMar',
             'tblLook', 'tblCaption', 'tblDescription']
    for table in doc.tables:
        tblpr = table._tbl.tblPr
        kids = list(tblpr)
        kids.sort(key=lambda e: order.index(e.tag.split('}')[1]) if e.tag.split('}')[1] in order else 99)
        for k in kids:
            tblpr.remove(k)
        for k in kids:
            tblpr.append(k)

    # keep small tables on one page; leave room after every table
    from docx.text.paragraph import Paragraph
    for table in doc.tables:
        nrows = len(table.rows)
        if nrows <= 24:
            for row in table.rows[:-1]:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        para.paragraph_format.keep_with_next = True
        nxt = table._tbl.getnext()
        if nxt is not None and nxt.tag == qn('w:p'):
            Paragraph(nxt, doc.element.body).paragraph_format.space_before = Pt(10)

    # References follows the conclusion on the same page
    for para in doc.paragraphs:
        if para.style.name == 'Heading 1' and para.text.strip() == 'References':
            para.paragraph_format.page_break_before = False
            para.paragraph_format.space_before = Pt(24)

    # table of contents
    headings = []
    for p in doc.paragraphs:
        if p.style.name in ('Heading 1', 'Heading 2'):
            headings.append((1 if p.style.name == 'Heading 1' else 2, p.text.strip()))
    marker = None
    for p in doc.paragraphs:
        if p.text.strip() == 'TOCMARKER':
            marker = p
            break
    if marker is not None:
        anchor = marker
        for level, text in headings:
            newp = OxmlElement('w:p')
            anchor._p.addnext(newp)
            from docx.text.paragraph import Paragraph
            para = Paragraph(newp, anchor._parent)
            para.style = doc.styles['Body Text']
            pf = para.paragraph_format
            pf.left_indent = Cm(0 if level == 1 else 0.8)
            pf.space_after = Pt(3 if level == 2 else 4)
            pf.space_before = Pt(6 if level == 1 else 0)
            pf.line_spacing = 1.0
            pf.tab_stops.add_tab_stop(Cm(TEXT_W_CM), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
            r = para.add_run(text.replace('\t', '  ')); r.font.name = FONT; r.font.size = Pt(12 if level == 1 else 11); r.font.bold = (level == 1)
            pg = '' if pages is None else str(pages.get(text, ''))
            r2 = para.add_run('\t' + pg); r2.font.name = FONT; r2.font.size = Pt(12 if level == 1 else 11)
            anchor = para
        marker._p.getparent().remove(marker._p)

    cp = doc.core_properties
    cp.title = 'MATLAB-Based Public-Key Cryptography and Secure Communication'
    cp.author = 'Santanu Barua, M Hariharan, Kushi Pathak, P Abdul Wasim, Someshwar A, Adithya Sampath Kumar'
    cp.subject = 'Micro-project report, Alliance University'
    doc.save(path_out)
    return headings


def find_pages(pdf, headings):
    n = int(re.search(r'Pages:\s+(\d+)', run(['pdfinfo', pdf]).stdout).group(1))
    texts = []
    for i in range(1, n + 1):
        t = run(['pdftotext', '-f', str(i), '-l', str(i), '-layout', pdf, '-']).stdout
        texts.append([ln.strip() for ln in t.splitlines()])
    pages = {}
    start = 2  # skip the cover and the contents page (page index 0 and 1)
    cur = start
    for level, text in headings:
        key = re.sub(r'\s+', ' ', text)
        for i in range(cur, n):
            if any(re.sub(r'\s+', ' ', ln) == key or (len(key) > 30 and re.sub(r'\s+', ' ', ln).startswith(key[:30])) for ln in texts[i]):
                pages[text] = i + 1
                cur = i
                break
    return pages, n


def to_pdf(docx, outdir):
    wrapper = '/mnt/skills/public/docx/scripts/office/soffice.py'
    cmd = ['python3', wrapper] if os.path.exists(wrapper) else ['soffice']
    run(cmd + ['--headless', '--convert-to', 'pdf', '--outdir', outdir, docx])
    return os.path.join(outdir, os.path.splitext(os.path.basename(docx))[0] + '.pdf')


def main():
    tmp = tempfile.mkdtemp()
    raw = os.path.join(tmp, 'raw.docx')
    run(['pandoc', SRC, '-f', 'markdown+smart', '-o', raw, '--number-sections', '--resource-path=' + ROOT], cwd=ROOT)
    p1 = os.path.join(tmp, 'pass1.docx')
    headings = style_document(raw, p1)
    pdf1 = to_pdf(p1, tmp)
    pages, n = find_pages(pdf1, headings)
    missing = [t for _, t in headings if t not in pages]
    if missing:
        print('WARNING: page not found for:', missing)
    style_document(raw, OUT_DOCX, pages)
    pdf2 = to_pdf(OUT_DOCX, os.path.dirname(OUT_DOCX))
    print('built', OUT_DOCX, 'and', pdf2, '(%d pages)' % n)
    json.dump(pages, open(os.path.join(tmp, 'pages.json'), 'w'), indent=1)
    shutil.rmtree(tmp, ignore_errors=True)


if __name__ == '__main__':
    main()
